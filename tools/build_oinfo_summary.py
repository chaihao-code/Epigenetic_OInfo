#!/usr/bin/env python3
"""Build readable full-system and triplet tables from existing result files."""
from pathlib import Path
import argparse
import html
import json
import re

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'supplementary_data/oinfo_summary'
FINAL = ROOT / 'supplementary_data/final_tables'
ALPHA = .05
CHECKS = {'matched_full_system_permutation_rows': 0,
          'maximum_observed_omega_difference': 0.,
          'maximum_tc_minus_dtc_difference': 0.,
          'bcell_duplicate_full_system_rows_verified': 0}


def rel(path):
    return str(path.relative_to(ROOT))


def bh(values):
    p = np.asarray(values, dtype=float)
    assert np.isfinite(p).all() and ((0 <= p) & (p <= 1)).all()
    order = np.argsort(p, kind='stable')
    q = np.minimum.accumulate((p[order]*len(p)/np.arange(1,len(p)+1))[::-1])[::-1]
    result = np.empty_like(q); result[order] = np.minimum(q, 1.)
    return result


def observed(path):
    columns = ['system','block','order','n_missing','remaining_names','hoi','tc','dtc']
    frame = pd.read_csv(path, usecols=columns)
    frame = frame.loc[frame.n_missing.eq(0)].copy()
    assert len(frame) and not frame.duplicated(['system','block']).any(), path
    error = (frame.hoi-(frame.tc-frame.dtc)).abs().max()
    assert error < 1e-10, (path,error)
    CHECKS['maximum_tc_minus_dtc_difference'] = max(CHECKS['maximum_tc_minus_dtc_difference'],float(error))
    return frame


def full_system():
    parts = []
    runs = [
        ('mouse_preimp','cohort_reference','cases/mouse_preimp/synergy_redundancy/results_n_bins_5'),
        ('mouse_development','triplet_perm10000_run','cases/mouse_development/result_oinfo_n_bins_5_triplet_perm10000'),
        ('mouse_after_E11.5','cohort_reference','cases/mouse_after_E11.5/results_oinfo_n_bins_5'),
        ('bcells_and_CLL','cohort_reference','cases/bcells_and_CLL/results_oinfo_n_bins_5'),
        ('aging/neutrophil','cohort_reference','cases/aging/neutrophil/results_oinfo_n_bins_5'),
        ('thyroid_demo','from_bins_current_inputs','demo/results/from_bins/oinfo'),
    ]
    for case,run,folder in runs:
        paths = sorted(p for p in (ROOT/folder).rglob('*_oinfo.csv') if not p.name.endswith('_permutation_oinfo.csv'))
        assert paths, folder
        for path in paths:
            sample = path.name.removesuffix('_oinfo.csv')
            data = observed(path)
            data['case']=case; data['run']=run; data['sample']=sample
            data['source_oinfo']=rel(path)
            summary = path.with_name(sample+'_permutation_summary.csv')
            if summary.exists():
                stats = pd.read_csv(summary)
                assert not stats.duplicated(['system','block']).any()
                data = data.merge(stats,on=['system','block'],how='left',validate='one_to_one',indicator=True)
                assert data._merge.eq('both').all() and len(data)==len(stats), summary
                error = (data.hoi-data.observed_hoi).abs().max()
                assert error < 1e-10, (summary,error)
                CHECKS['matched_full_system_permutation_rows'] += len(data)
                CHECKS['maximum_observed_omega_difference'] = max(CHECKS['maximum_observed_omega_difference'],float(error))
                data['q_two_sided_bh']=np.nan
                data['fdr_family_size']=0
                for _,indexes in data.groupby('system').groups.items():
                    # New table-level correction: all 40 regions in this sample/system.
                    assert len(indexes)==40
                    data.loc[indexes,'q_two_sided_bh']=bh(data.loc[indexes,'p_dev'])
                    data.loc[indexes,'fdr_family_size']=len(indexes)
                data['fdr_family']='within_run_sample_system_all_40_regions'
                data['significance_status']='tested'
                data['source_significance']=rel(summary)
            else:
                assert case=='mouse_development', (case,path)
                data['significance_status']='full_system_not_tested_triplet_null_only'
            if case=='bcells_and_CLL':
                duplicate = ROOT/'cases/bcells_and_CLL/results_oinfo_triplet_n_bins_5'/sample/path.name
                if duplicate.exists():
                    other=observed(duplicate)
                    comparison=data.merge(other,on=['system','block'],suffixes=('','_triplet'),validate='one_to_one')
                    assert len(comparison)==len(data)
                    assert (comparison.hoi-comparison.hoi_triplet).abs().max()<1e-10
                    CHECKS['bcell_duplicate_full_system_rows_verified']+=len(data)
            parts.append(data)
    thyroid = ROOT/'cases/aging/thyroid/deconvolution/data/processed/oinfo_long.tsv'
    data=pd.read_csv(thyroid,sep='\t').rename(columns={'sample_id':'sample','region':'block'})
    error=(data.hoi-(data.tc-data.dtc)).abs().max()
    assert error<1e-10
    CHECKS['maximum_tc_minus_dtc_difference']=max(CHECKS['maximum_tc_minus_dtc_difference'],float(error))
    data['case']='aging/thyroid';data['run']='historical_six_donor_table';data['order']=7
    data['significance_status']='historical_permutation_results_unavailable'
    data['source_oinfo']=rel(thyroid)
    parts.append(data)
    result=pd.concat(parts,ignore_index=True).rename(columns={
        'block':'region','order':'n_signals','hoi':'oinfo_bits','system':'legacy_system_label',
        'p_dev':'p_two_sided','p_down':'p_lower','p_up':'p_upper','n_perm':'n_permutations'})
    result['omega_sign']=np.select([result.oinfo_bits.lt(0),result.oinfo_bits.gt(0)],['negative','positive'],default='zero')
    result['n_permutations']=result.n_permutations.astype('Int64')
    result['fdr_family_size']=result.fdr_family_size.astype('Int64')
    result['statistical_call']='not_tested'
    tested=result.significance_status.eq('tested')
    result.loc[tested,'statistical_call']='not_significant'
    hit=tested & result.q_two_sided_bh.le(ALPHA)
    result.loc[hit & result.oinfo_bits.gt(result.null_mean),'statistical_call']='significantly_above_permutation_null'
    result.loc[hit & result.oinfo_bits.lt(result.null_mean),'statistical_call']='significantly_below_permutation_null'
    assert not result.duplicated(['case','run','sample','legacy_system_label','region']).any()
    columns=['case','run','sample','region','n_signals','oinfo_bits','p_two_sided','q_two_sided_bh',
             'statistical_call','significance_status','n_permutations','omega_sign','p_lower','p_upper',
             'null_mean','null_sd','tc','dtc','fdr_family','fdr_family_size','remaining_names',
             'legacy_system_label','source_oinfo','source_significance']
    return result.reindex(columns=columns).sort_values(['case','run','sample','region']).reset_index(drop=True)


def triplets():
    parts=[]
    configs=[
        ('mouse_development','cases/mouse_development/synergistic_triplets/all_triplets_fdr_10000.tsv.gz',
         'q_down_sample_block_bh','sample_region_BH35',35),
        ('bcells_and_CLL','cases/bcells_and_CLL/results_oinfo_triplet_n_bins_5/all_samples_triplet_significance.tsv',
         'q_down_global_BH','all_17_samples_40_regions_35_triplets_BH23800',23800),
    ]
    for case,source,qcol,family,size in configs:
        data=pd.read_csv(ROOT/source,sep='\t')
        assert data.order.eq(3).all() and len(data)
        groups=data.groupby(['sample','block']) if case=='mouse_development' else [(None,data)]
        max_q_error=0.
        for _,g in groups:
            assert len(g)==size
            err=np.max(np.abs(bh(g.p_down)-g[qcol]))
            assert err<1e-10,(case,err)
            max_q_error=max(max_q_error,float(err))
        CHECKS[case+'_triplet_maximum_saved_bh_difference']=max_q_error
        data['case']=case;data['q_lower_bh']=data[qcol]
        data['fdr_family']=family;data['fdr_family_size']=size;data['source_significance']=source
        data=data.rename(columns={'block':'region','subset_names':'signals','observed_hoi':'oinfo_bits',
                                 'p_down':'p_lower','p_up':'p_upper','p_dev':'p_two_sided','n_perm':'n_permutations'})
        def short_signals(value):
            marks=[]
            for name in value.split('/'):
                match=re.search(r'H3K\d+(?:me[123]|ac)|ATAC|WGBS|WGB-Seq|CpG_methyl',name)
                assert match,(case,name)
                mark=match.group()
                marks.append('CpG_methyl' if mark in ('WGBS','WGB-Seq') else mark)
            assert len(marks)==len(set(marks))==3
            return '/'.join(sorted(marks))
        data['triplet_marks']=data['triplet'] if 'triplet' in data else data.signals.map(short_signals)
        assert not data.duplicated(['sample','region','signals']).any()
        data['significant_negative_omega_lower_tail']=data.oinfo_bits.lt(0)&data.q_lower_bh.le(ALPHA)
        if 'significant_negative' in data:
            assert data.significant_negative.eq(data.significant_negative_omega_lower_tail).all()
        parts.append(data)
    cols=['case','sample','region','triplet_marks','signals','oinfo_bits','p_lower','q_lower_bh',
          'significant_negative_omega_lower_tail','n_permutations','fdr_family','fdr_family_size',
          'p_upper','p_two_sided','null_mean','null_sd','q_down_sample_BH','q_down_global_BH',
          'q_dev_sample_BH','q_dev_global_BH','source_significance']
    return pd.concat(parts,ignore_index=True).reindex(columns=cols).sort_values(['case','sample','region','signals']).reset_index(drop=True)


CALLS={'significantly_above_permutation_null':'Significantly above permutation null',
       'significantly_below_permutation_null':'Significantly below permutation null',
       'not_significant':'Not significant','not_tested':'Not tested / unavailable'}


def excel_sheet(book,name,frame,widths=None):
    if book.loaded_theme is None:
        # Office's default theme embeds localized font names even in English
        # workbooks. Keep the entire exported workbook free of non-English text.
        from openpyxl.writer.theme import theme_xml
        theme=re.sub(r'typeface="([^"]*)"',
                     lambda m:m.group(0) if html.unescape(m.group(1)).isascii() else 'typeface="Arial"',theme_xml)
        book.loaded_theme=theme.encode('utf-8')
    sheet=book.create_sheet(name);sheet.freeze_panes='A2'
    sheet.auto_filter.ref=f'A1:{get_column_letter(len(frame.columns))}{len(frame)+1}'
    for i,c in enumerate(frame.columns,1):
        sheet.column_dimensions[get_column_letter(i)].width=(widths or {}).get(c,22)
    header=[]
    for label in frame.columns:
        cell=WriteOnlyCell(sheet,value=label);cell.font=Font(bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='315A7D');header.append(cell)
    sheet.append(header)
    for row in frame.itertuples(index=False,name=None):
        sheet.append([None if pd.isna(v) else v for v in row])


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--reuse-tables',action='store_true',help='Export the workbook from the existing verified TSV tables')
    args=parser.parse_args()
    OUT.mkdir(parents=True,exist_ok=True)
    FINAL.mkdir(parents=True,exist_ok=True)
    if args.reuse_tables:
        full=pd.read_csv(OUT/'full_system_oinfo.tsv',sep='\t',float_precision='round_trip')
        trip=pd.read_csv(OUT/'triplet_oinfo.tsv.gz',sep='\t',float_precision='round_trip')
        CHECKS.update(json.loads((OUT/'checks.json').read_text()))
        for key in ['xlsx_zip_crc_passed','xlsx_row_counts_including_header']:
            CHECKS.pop(key,None)
    else:
        full=full_system()
        trip=triplets()
    print('Full-system rows:',len(full),flush=True)
    print('Triplet rows:',len(trip),flush=True)
    full.to_csv(OUT/'full_system_oinfo.tsv',sep='\t',index=False,na_rep='NA')
    trip.to_csv(OUT/'triplet_oinfo.tsv.gz',sep='\t',index=False,na_rep='NA',compression='gzip')
    brief=full[['case','sample','region','n_signals','oinfo_bits','p_two_sided','q_two_sided_bh',
                'statistical_call','n_permutations','run']].copy()
    brief.statistical_call=brief.statistical_call.map(CALLS)
    brief=brief.rename(columns={'case':'Case','sample':'Sample','region':'Genome_region','n_signals':'N_signals',
         'oinfo_bits':'O_information_bits','p_two_sided':'Permutation_P_two_sided','q_two_sided_bh':'BH_q_40_regions',
         'statistical_call':'Significance_q_le_0.05','n_permutations':'N_permutations','run':'Result_version'})
    coverage=[]
    for (case,run),g in full.groupby(['case','run'],sort=True):
        coverage.append(dict(case=case,run=run,samples=g['sample'].nunique(),rows=len(g),
             with_permutation_p=int(g.p_two_sided.notna().sum()),without_permutation_p=int(g.p_two_sided.isna().sum()),
             significant_bh05=int(g.q_two_sided_bh.le(ALPHA).sum()) if g.p_two_sided.notna().any() else None,
             n_signals='/'.join(map(str,sorted(g.n_signals.unique())))))
    coverage=pd.DataFrame(coverage);coverage.to_csv(OUT/'coverage.tsv',sep='\t',index=False,na_rep='NA')
    guide=[
        ('Full system','Each row represents one case, result version, sample and annotation region using all signals. O-information is measured in bits.'),
        ('Permutation P values','Read from the matching permutation_summary file; observed_hoi was checked against the original observed result for every tested row.'),
        ('Full-system q values','New table-level Benjamini-Hochberg correction of two-sided p_dev across all 40 regions within each result version, sample and system. Permutations were not rerun.'),
        ('Full-system significance','At q <= 0.05, direction is determined relative to null_mean. The sign of O-information alone does not establish a significant null deviation.'),
        ('Missing values','Blank / NA means not tested or no matching statistical result. It does not mean not significant.'),
        ('Mouse development','The retained permutation run tested triplets only. Seven-signal full-system O-information is available, but full-system P and q values are not.'),
        ('Historical thyroid cohort','No matching permutation summaries are available for the six-donor historical table. P and q values remain blank.'),
        ('Thyroid demo','The current-input from_bins run uses 10000 permutations and is listed separately from the historical cohort. Equivalent from_bigwig results are not duplicated.'),
        ('Triplet q values','Original lower-tail BH q values are retained: mouse sample-by-region BH35; B-cell global BH23800 across 17 samples, 40 regions and 35 triplets.'),
        ('Significant negative triplets','Both O-information < 0 and lower-tail q <= 0.05 are required. Multiple-testing correction includes all triplet tests before filtering by sign.'),
        ('Legacy system labels','Historical 9sig labels may describe 7 or 10 actual signals. Signal counts use order / remaining_names, not the legacy label.'),
        ('Scope','The full-system table covers six original cohorts and one complete demo. Triplet sheets cover retained FDR results for mouse_development and B cells.'),
        ('Other subsystem orders','Observed results for other orders remain in cases/*/tables/per_sample. They are outside the scope of this compact workbook.'),
        ('H3K36me3 sensitivity branches','The 16 source-replacement variants and shared-eight-signal comparisons are separate sensitivity analyses and are outside this workbook.'),
        ('Reproduction','Run python tools/build_oinfo_summary.py. Detailed original result paths are provided in the source_* columns of the TSV exports.'),
    ]
    book=Workbook(write_only=True)
    excel_sheet(book,'Readme',pd.DataFrame(guide,columns=['Item','Description']),{'Item':30,'Description':110})
    excel_sheet(book,'Full_system',brief,{'Genome_region':64,'Case':25,'Sample':26,'Result_version':30,'Significance_q_le_0.05':45})
    for case,title in [('mouse_development','Triplets_mouse_development'),('bcells_and_CLL','Triplets_B_cells')]:
        frame=trip[trip.case.eq(case)][['sample','region','triplet_marks','oinfo_bits','p_lower','q_lower_bh',
                  'significant_negative_omega_lower_tail','n_permutations','fdr_family']].copy()
        frame=frame.rename(columns={'sample':'Sample','region':'Genome_region','triplet_marks':'Triplet_signals','oinfo_bits':'O_information_bits',
              'p_lower':'Permutation_P_lower_tail','q_lower_bh':'Original_BH_q_lower_tail','significant_negative_omega_lower_tail':'Negative_Omega_and_q_le_0.05',
              'n_permutations':'N_permutations','fdr_family':'Multiple_testing_family'})
        excel_sheet(book,title,frame,{'Genome_region':64,'Triplet_signals':45,'Multiple_testing_family':58})
    excel_sheet(book,'Coverage',coverage)
    book.save(FINAL/'O_information_summary.xlsx')
    CHECKS.update(full_system_rows=len(full),triplet_rows=len(trip),
                  full_system_rows_with_p=int(full.p_two_sided.notna().sum()),
                  full_system_rows_without_p=int(full.p_two_sided.isna().sum()),
                  excel_sheets=book.sheetnames,alpha=ALPHA)
    (OUT/'checks.json').write_text(json.dumps(CHECKS,ensure_ascii=False,indent=2)+'\n')
    lines=['# O-information summary tables','',
           f'{len(full):,} full-system rows and {len(trip):,} triplet rows. Existing results are summarized without rerunning O-information or permutations.','',
           'Open the [English Excel workbook](../final_tables/O_information_summary.xlsx): Readme, Full_system, two triplet sheets and Coverage.','',
           '| Case | Samples | Full-system rows | With P values | Without P values |','|---|---:|---:|---:|---:|']
    for r in coverage.itertuples(index=False):lines.append(f'| {r.case} ({r.run}) | {r.samples} | {r.rows} | {r.with_permutation_p} | {r.without_permutation_p} |')
    lines+=['','## Fields and statistical interpretation','']+[f'- **{key}**: {value}' for key,value in guide]
    lines+=['','## Files','',
            '- `full_system_oinfo.tsv`: O-information, TC/DTC, P/q values, exact signal names, missing-test reasons and original source paths.',
            '- `triplet_oinfo.tsv.gz`: All retained triplet tests, with original P/q values and explicit multiple-testing families.',
            '- `coverage.tsv`: Coverage by case and result version.',
            '- `checks.json`: Observed-value matching, TC minus DTC identity checks and verification of original BH results.','',
            'Full-system q values use the new BH40 correction defined above. Triplet q values retain their original methods and testing families.']
    (OUT/'README.md').write_text('\n'.join(lines)+'\n')
    print(json.dumps(CHECKS,ensure_ascii=False),flush=True)


if __name__=='__main__':
    main()
